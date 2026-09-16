import os
import time
import httpx
import docx
from pypdf import PdfWriter

BASE_URL = "http://localhost:8000"
NOTEBOOK_ID = "default_notebook_1"


def create_sample_files():
    """Generate sample TXT, MD, DOCX, and PDF test files."""
    os.makedirs("test_files", exist_ok=True)

    # 1. TXT File
    txt_path = "test_files/sample_research.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(
            "Artificial Intelligence & Machine Learning Overview\n\n"
            "Artificial intelligence (AI) is intelligence demonstrated by machines, as opposed to intelligence displayed by animals including humans.\n\n"
            "Machine learning (ML) is a field of study in artificial intelligence concerned with the development and study of statistical algorithms that can learn from data and generalize to unseen data, and thus perform tasks without explicit instructions."
        )

    # 2. Markdown File
    md_path = "test_files/sample_guide.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(
            "# Next.js and FastAPI RAG Architecture\n\n"
            "## Introduction\n\n"
            "Retrieval-Augmented Generation (RAG) is a technique for enhancing AI model responses with authoritative information retrieved from external knowledge bases.\n\n"
            "```python\n"
            "def query_rag(prompt):\n"
            "    return 'RAG response'\n"
            "```\n\n"
            "### Key Benefits\n"
            "- Grounded answers\n"
            "- Reduced hallucinations\n"
            "- Domain-specific intelligence"
        )

    # 3. DOCX File
    docx_path = "test_files/sample_contract.docx"
    doc = docx.Document()
    doc.add_heading("Software Engineering Agreement", level=1)
    doc.add_paragraph(
        "This Software Engineering Agreement is entered into by and between Recall AI Platform and the Client. "
        "The developer agrees to deliver high quality modular code built with Next.js, FastAPI, and MongoDB."
    )
    doc.save(docx_path)

    # 4. PDF File
    pdf_path = "test_files/sample_paper.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    # 5. Invalid File (.exe)
    invalid_path = "test_files/invalid_file.exe"
    with open(invalid_path, "wb") as f:
        f.write(b"\x4d\x5a\x90\x00\x03\x00\x00\x00")

    return [txt_path, md_path, docx_path, pdf_path], invalid_path


def run_tests():
    print("--- Starting Step 2 End-to-End Document Upload & Processing Test ---")
    valid_files, invalid_file = create_sample_files()

    client = httpx.Client(base_url=BASE_URL, timeout=10.0)

    # Test 1: Upload valid files
    uploaded_source_ids = []
    for file_path in valid_files:
        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            files = {"file": (filename, f.read())}
            res = client.post(f"/api/v1/notebooks/{NOTEBOOK_ID}/sources/upload", files=files)
            print(f"Upload {filename}: Status {res.status_code}")
            assert res.status_code == 202, f"Failed upload for {filename}: {res.text}"
            source_data = res.json()
            uploaded_source_ids.append(source_data["id"])
            print(f"  -> Created Source ID: {source_data['id']} (Status: {source_data['status']})")

    # Test 2: Upload invalid file type (.exe)
    with open(invalid_file, "rb") as f:
        files = {"file": (os.path.basename(invalid_file), f.read())}
        res = client.post(f"/api/v1/notebooks/{NOTEBOOK_ID}/sources/upload", files=files)
        print(f"Upload invalid file (.exe): Status {res.status_code}")
        assert res.status_code == 400, "Expected 400 Bad Request for invalid file type"
        print(f"  -> Successfully rejected invalid file: {res.json()['detail']}")

    # Wait for background processing tasks to complete
    print("Waiting 3 seconds for background document processing & chunking...")
    time.sleep(3)

    # Test 3: List sources
    res = client.get(f"/api/v1/notebooks/{NOTEBOOK_ID}/sources")
    print(f"List sources: Status {res.status_code}")
    assert res.status_code == 200
    sources_list = res.json()
    print(f"  -> Found {len(sources_list)} sources in notebook '{NOTEBOOK_ID}'")

    # Test 4: Check source detail & extracted chunks for first source
    if uploaded_source_ids:
        test_id = uploaded_source_ids[0]
        res = client.get(f"/api/v1/sources/{test_id}")
        print(f"Get source detail ({test_id}): Status {res.status_code}")
        assert res.status_code == 200
        detail = res.json()
        source = detail["source"]
        chunks = detail["chunks"]
        print(f"  -> Source '{source['title']}' Status: {source['status']}")
        print(f"  -> Extracted {len(chunks)} chunks, total {source['char_count']} characters.")
        if chunks:
            print(f"  -> Sample Chunk 0 text: '{chunks[0]['text'][:80]}...'")

    # Test 5: Delete one source
    if uploaded_source_ids:
        delete_id = uploaded_source_ids[-1]
        res = client.delete(f"/api/v1/notebooks/{NOTEBOOK_ID}/sources/{delete_id}")
        print(f"Delete source ({delete_id}): Status {res.status_code}")
        assert res.status_code == 200
        print(f"  -> Deleted source successfully!")

    print("\n✅ ALL STEP 2 DOCUMENT PROCESSING TESTS PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    run_tests()
