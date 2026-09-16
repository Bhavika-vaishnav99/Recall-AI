import io
import logging
from typing import List, Dict, Any
import fitz
import docx

from app.document_processing.text_cleaning import clean_text

logger = logging.getLogger(__name__)


def extract_text_from_file(file_bytes: bytes, file_type: str) -> List[Dict[str, Any]]:
    """
    Extract raw text from PDF, DOCX, TXT, or Markdown file bytes.
    Returns list of page items: [{"page": page_num, "text": "cleaned text"}, ...]
    """
    if file_type == "pdf":
        return extract_from_pdf(file_bytes)
    elif file_type == "docx":
        return extract_from_docx(file_bytes)
    elif file_type in ("txt", "markdown"):
        return extract_from_text_plain(file_bytes)
    else:
        raise ValueError(f"Unsupported document type: {file_type}")


def extract_from_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract page-by-page text from PDF using PyMuPDF (fitz)."""
    pages = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for idx, page in enumerate(doc):
            page_text = page.get_text() or ""
            cleaned = clean_text(page_text)
            if cleaned:
                pages.append({
                    "page": idx + 1,
                    "text": cleaned
                })
        doc.close()
    except Exception as exc:
        logger.error("Error extracting text from PDF: %s", exc)
        raise RuntimeError(f"Failed to parse PDF document: {str(exc)}") from exc

    return pages


def extract_from_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract text from Microsoft Word (.docx) document."""
    try:
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        
        full_text_paragraphs = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_text_paragraphs.append(paragraph.text)
                
        # Also extract table cells text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text_paragraphs.append(" | ".join(row_text))

        combined_text = "\n\n".join(full_text_paragraphs)
        cleaned = clean_text(combined_text)
        
        return [{"page": 1, "text": cleaned}] if cleaned else []
    except Exception as exc:
        logger.error("Error extracting text from DOCX: %s", exc)
        raise RuntimeError(f"Failed to parse Word document: {str(exc)}") from exc


def extract_from_text_plain(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract text from TXT or Markdown file trying multiple encodings."""
    raw_str = ""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            raw_str = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    cleaned = clean_text(raw_str)
    return [{"page": 1, "text": cleaned}] if cleaned else []
